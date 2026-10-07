# -*- coding: utf-8 -*-
"""Revise group-meeting-2026-10-06.pptx to content-outline v3 (8 pages),
in place, content-addressed.

v3 target order:
  1 封面 | 2 这几天工作总览 | 3 需求与信号链梳理 | 4 数据接口与协议初稿
  5 系统架构与产品方案 (NEW) | 6 团队分工 | 7 需要老师确认 | 8 下一步 (NEW)

Also fixes a real defect found while building v3: NO slide had an explicit
<p:bg>, so PowerPoint rendered the cover white-on-white and the content
pages pure white instead of the designed cream. Every slide now gets an
explicit solid background (cover INK, others F5F2EC).
"""
import copy
import json
import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt

import deckkit as dk
from deckkit import band, header, notes, rect, txt, img

ROOT = Path(__file__).resolve().parents[3]
MEETING_DIR = ROOT / "meetings" / "2026-10-06"
DECK = MEETING_DIR / "group-meeting-2026-10-06.pptx"
TMP = Path(os.environ.get("LOCALAPPDATA", r"C:\Temp")) / "Temp"
GROUPS_OUT = TMP / "deck_anim_groups.json"
ARCH = ROOT / "assets" / "fyp-system-architecture.png"
PRODUCT = TMP / "fyp-product-crop.png"

PNS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
BG_XML = ('<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
          '<p:bgPr><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
          '<a:effectLst/></p:bgPr></p:bg>')


# ---------------------------------------------------------------- helpers
def set_bg(slide, color):
    cSld = slide._element.find(PNS + "cSld")
    old = cSld.find(PNS + "bg")
    if old is not None:
        cSld.remove(old)
    cSld.insert(0, parse_xml(BG_XML % color))


def set_notes(slide, text, src_sp=None):
    """notes_slide.notes_text_frame is None on slides ADDED to a loaded
    package (fresh notes slides carry no body placeholder). Borrow a working
    body <p:sp> from an existing notes slide, bump its id, then set text."""
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        if src_sp is None:
            raise SystemExit("no notes source sp available")
        sp = copy.deepcopy(src_sp)
        cNvPr = sp.find(".//" + qn("p:cNvPr"))
        used = [int(e.get("id")) for e in
                slide.notes_slide._element.iter(qn("p:cNvPr"))]
        cNvPr.set("id", str(max(used) + 1 if used else 2))
        slide.notes_slide._element.spTree.append(sp)
        tf = slide.notes_slide.notes_text_frame
    tf.text = text


def find_slide(prs, *markers):
    """Find a slide by ANY of the markers, so re-running on an already
    revised deck still works (old title or new title)."""
    for marker in markers:
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame and marker in sh.text_frame.text:
                    return s
    raise SystemExit(f"marker not found: {markers}")


def wipe(slide):
    for el in list(slide._element):
        if el.tag == PNS + "timing":
            slide._element.remove(el)
    spTree = slide.shapes._spTree
    for sh in list(slide.shapes):
        spTree.remove(sh._element)


def replace_text(slide, marker, new):
    """Replace the whole paragraph text of the single shape containing marker."""
    for sh in slide.shapes:
        if sh.has_text_frame and marker in sh.text_frame.text:
            para = sh.text_frame.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
                return True
    return False


def section_label(slide, x, y, text, color, size=14.5):
    rect(slide, x, y + 0.04, 0.13, 0.24, fill=color)
    txt(slide, x + 0.24, y, 5.0, 0.32,
        [{"text": text, "size": size, "bold": True, "color": color}])


def bullet_block(slide, x, y, w, h, items, size=16, gap=7, line_spacing=1.25):
    paras = [{"text": "•  " + it, "size": size, "color": dk.BODY,
              "line_spacing": line_spacing, "space_after": gap} for it in items]
    return txt(slide, x, y, w, h, paras)


ROW_TINT = "F1EBE1"
CONFIRM = [
    ("1 · 设备结构", dk.CORAL, [
        "已确定采用二体式：探头贴胸采集，主机独立完成显示、控制和联网",
        "请老师确认这个总体结构是否合适",
        "请老师确认探头与主机之间采用 BLE 无线还是有线连接",
        "请老师确认主机的基本形态是手持还是随身佩戴",
    ]),
    ("2 · 上传服务器的范围", dk.AMBER, [
        "服务器只接收 AI 结果和异常告警，还是也需要特征、BPM 和信号质量",
        "是否需要上传原始心音或波形，供服务器回看和生成报告",
        "测量过程中实时上传，还是测量结束后按会话批量上传",
        "AI 推理固定在本地主机运行，服务器不承担模型推理",
    ]),
]


def confirm_card(slide, cx, w, accent, heading, rows):
    y, h = dk.CONTENT_TOP, 4.35
    rect(slide, cx, y, w, h, fill=dk.WHITE, line_color=dk.CARD_LINE,
         rounded=True, radius=0.05)
    rect(slide, cx, y, w, 0.6, fill=accent, rounded=True, radius=0.07)
    rect(slide, cx, y + 0.46, w, 0.14, fill=accent)
    rect(slide, cx, y, w, 0.6,
         text=[{"text": heading, "size": 21, "bold": True, "color": dk.WHITE,
                "align": "l"}],
         anchor="MIDDLE", pad=(0.3, 0.02, 0.1, 0.02))
    for n, body in enumerate(rows):
        rect(slide, cx + 0.3, 2.71 + n * 0.86, w - 0.6, 0.74, fill=ROW_TINT,
             rounded=True, radius=0.12,
             text=[{"text": "%d.   %s" % (n + 1, body), "size": 16,
                    "color": dk.INK, "line_spacing": 1.22}],
             anchor="MIDDLE", pad=(0.24, 0.04, 0.16, 0.04))


def white_card(slide, x, y, w, h):
    return rect(slide, x, y, w, h, fill=dk.WHITE, line_color=dk.CARD_LINE,
                rounded=True, radius=0.05)


def numbered_row(slide, x, y, w, h, n, text, accent, size=16):
    rect(slide, x, y, w, h, fill=dk.WHITE, line_color=dk.CARD_LINE,
         rounded=True, radius=0.14)
    rect(slide, x + 0.22, y + h / 2 - 0.16, 0.32, 0.32, fill=accent,
         rounded=True, radius=0.5,
         text=[{"text": str(n), "size": 14, "bold": True, "color": dk.WHITE,
                "align": "c"}],
         anchor="MIDDLE", pad=(0, 0, 0, 0))
    txt(slide, x + 0.68, y, w - 0.9, h,
        [{"text": text, "size": size, "color": dk.INK}],
        anchor="MIDDLE")


def output_row(slide, x, y, w, h, text, accent, size=16):
    rect(slide, x, y, w, h, fill=dk.WHITE, line_color=dk.CARD_LINE,
         rounded=True, radius=0.14)
    rect(slide, x, y + 0.12, 0.1, h - 0.24, fill=accent)
    txt(slide, x + 0.3, y, w - 0.5, h,
        [{"text": text, "size": size, "color": dk.INK}], anchor="MIDDLE")


FLOW = [
    dict(x=0.9, w=3.05, fill=dk.TEAL, line=None, fg=dk.WHITE,
         text="探头（贴胸口）", size=17),
    dict(x=5.1, w=3.55, fill=dk.INK, line=None, fg=dk.WHITE,
         text="主机（AI + 屏幕 + 扬声器 + WiFi）", size=15),
    dict(x=9.8, w=2.63, fill=dk.WHITE, line=dk.AMBER, fg=dk.AMBER,
         text="服务器（数据存储）", size=17),
]
GAPS = [(3.95, 1.15, "BLE /  有线"), (8.65, 1.15, "WiFi / 手机网关")]


def flow_strip(s, top=1.85, height=0.9):
    for f in FLOW:
        rect(s, f["x"], top, f["w"], height, fill=f["fill"],
             line_color=f["line"], line_w=1.75, rounded=True, radius=0.1,
             text=[{"text": f["text"], "size": f["size"], "bold": True,
                    "color": f["fg"], "align": "c"}],
             anchor="MIDDLE", pad=(0.08, 0.04, 0.08, 0.04))
    for gx, gw, label in GAPS:
        txt(s, gx, top + 0.02, gw, 0.28,
            [{"text": label, "size": 11, "color": dk.MUTED, "align": "c"}])
        rect(s, gx + gw / 2 - 0.27, top + 0.37, 0.54, 0.26, fill=dk.MUTED,
             shape=MSO_SHAPE.RIGHT_ARROW)


# ---------------------------------------------------------------- content
OVERVIEW_WORK = [
    "梳理项目需求和老师给出的信号链",
    "完成阶段 0 的数据与硬件接口初稿",
    "明确探头、主机和服务器的总体关系",
    "制作系统架构图和产品交互示意图",
    "与队友确定分工，整理待确认的结构与上传范围",
]
OVERVIEW_OUT = [
    "阶段 0 接口冻结文档",
    "固件与上位机共用的音频协议头文件",
    "系统架构图",
    "产品外观和交互示意图",
]

NEEDS = [
    "项目目标是便携式数字蓝牙听诊器：心音采集、显示、播放、无线传输和边缘端初步诊断",
    "设备必须在脱离服务器时完成采集、显示、播放和本地 AI 推理",
    "服务器只作为外部数据存储，不参与实时诊断",
]
NOTES_TXT = [
    "有效心音带宽暂按 20 到 1000 Hz 处理",
    "20 到 1000 Hz 是信号带宽，不是采样率",
    "传感器候选：piezo/PVDF、麦克风、MEMS，型号待硬件确认",
    "先打通采集、显示、播放和传输，再逐步优化 AI 模型",
]
CHAIN = [
    ("传感器", 1.15, dk.TEAL), ("模拟前端", 1.35, dk.TEAL),
    ("ADC", 1.05, dk.TEAL), ("MCU 采集", 1.4, dk.TEAL),
    ("DSP 预处理", 1.55, dk.TEAL), ("本地 AI 推理", 1.7, dk.INK),
    ("显示·播放·无线发送", 2.5, dk.INK),
]

IF_LEFT = [
    "采集基线：8 kHz、单声道、16-bit PCM、小端",
    "AI 输入：带限后 2:1 抽取到 4 kHz",
    "逻辑帧：20 ms 一帧，8 kHz 下为 160 samples",
    "帧格式：320 B payload + 30 B header + 2 B CRC，共 352 B",
]
IF_RIGHT = [
    "每帧包含 sequence 和 timestamp_ms",
    "timestamp_ms 表示该帧第一个样本的采集时刻",
    "探头生成序号和时间戳，主机必须原样使用、不重新生成",
    "已提供固件和上位机共用的协议头文件",
    "已定义固定测试帧，用于双方核对 header 和 CRC",
]

ARCH_BULLETS = [
    "探头负责采集和预处理，不做 AI、不显示、不播放",
    "主机负责接收、显示、播放、本地 AI 推理、WiFi 上传和断网缓存",
    "服务器只负责数据存储，当前不包含训练、模型管理和复杂 GUI",
]
PRODUCT_BULLETS = [
    "架构图已制作：覆盖采集、组帧、传输、主机处理与结果展示",
    "交互示意图已制作：心音波形、BPM、AI 状态、信号质量",
    "尚未冻结探头与主板结构，未完成真实硬件联调",
]

STEPS = [
    ("1", "冻结结构与连接", "根据老师意见冻结二体式结构和连接方式", dk.CORAL),
    ("2", "确定上传范围", "确定服务器上传的数据范围", dk.AMBER),
    ("3", "细化器件与接口", "之后再与队友确认具体器件、接口和上传字段", dk.TEAL),
]

NOTES = {
    0: "这次汇报 6 到 8 分钟：第 2 到第 5 页展开讲这几天做了什么，"
       "后面是分工、需要老师确认的两件事和下一步。",
    1: "这几天我们主要做了三件事。第一，梳理了系统整体方案，把心音采集、数据传输、"
       "主机分析和结果显示这条链路理清了。第二，整理了设备端的数据接口初稿，"
       "确定 8 kHz 采集、20 ms 一帧，并加入序号、时间戳和 CRC 校验，"
       "方便后面检查丢帧和数据损坏。第三，制作了系统架构图和产品示意图，"
       "同时开始整理硬件与软件之间的接口。产出方面，阶段 0 接口冻结文档、"
       "共用的音频协议头文件、架构图和产品示意图都已经落盘。",
    2: "需求上，项目目标是便携式数字蓝牙听诊器，覆盖采集、显示、播放、无线传输和"
       "边缘端初步诊断；设备脱离服务器也要能完成本地 AI 推理，服务器只做外部存储。"
       "信号链是传感器、模拟前端、ADC、MCU 采集、DSP 预处理、本地 AI 推理，"
       "再到显示、播放和无线发送。有效心音带宽暂按 20 到 1000 Hz 处理，"
       "注意这是信号带宽不是采样率；传感器候选有 piezo/PVDF、麦克风和 MEMS，"
       "最终型号还要等硬件确认。策略是先打通链路，再逐步优化模型。",
    3: "第二件事是数据接口初稿：8 kHz、单声道、16-bit PCM、小端；AI 输入带限后"
       "2:1 抽取到 4 kHz；20 ms 一帧，8 kHz 下 160 个样本；"
       "一帧是 320 字节 payload 加 30 字节 header 加 2 字节 CRC，共 352 字节。"
       "每帧带序号和时间戳，timestamp_ms 表示该帧第一个样本的采集时刻，"
       "探头生成、主机原样使用。固件和上位机共用的协议头文件已经提供，"
       "固定测试帧用来双方核对 header 和 CRC。",
    4: "系统分成三部分：探头只负责采集和预处理；主机负责接收、显示、播放、"
       "本地 AI 推理、WiFi 上传和断网缓存；服务器只负责数据存储。"
       "已确定二体式，探头与主机分离。产出上，架构图覆盖采集、组帧、传输、"
       "主机处理和结果展示；交互示意图包含心音波形、BPM、AI 状态和信号质量。"
       "当前阶段还没有冻结探头和主板结构，也没有完成真实硬件联调。",
    5: "分工方面，我主要负责 AI 模型训练与边缘部署、显示、无线通信和软件集成；"
       "队友负责心音采集、模拟前端、回放、电源、PCB 和硬件结构。"
       "我们会共同冻结接口，并完成联调。",
    6: "目前已经确定采用二体式，探头贴胸采集，主机独立负责显示、控制、播放和联网。"
       "想请老师确认这个总体结构是否合适，以及探头和主机之间采用 BLE 无线还是有线连接、"
       "蓝牙是否为硬性要求，还有主机形态是手持还是随身佩戴。"
       "服务器方面，我们计划只保留数据存储功能，不参与 AI 推理、模型训练和模型管理；"
       "想请老师确认上传的数据范围，以及实时上传还是测量结束后上传。",
    7: "确认之后，我们再和队友细化器件、接口和上传字段：先冻结二体式结构和连接方式，"
       "再确定服务器上传的数据范围，最后一起确认具体器件、接口和上传字段。",
}


# ---------------------------------------------------------------- build
def build():
    prs = Presentation(DECK)
    s_over = find_slide(prs, "这几天做了什么", "这几天工作总览")
    s_arch = find_slide(prs, "系统总体架构", "需求与信号链梳理")
    s_duty = find_slide(prs, "三部分职责", "数据接口与协议初稿")
    s_team = find_slide(prs, "我负责")
    s_conf = find_slide(prs, "需要老师确认")

    # notes source: a body <p:sp> that already works on this package
    src_sp = None
    for s in prs.slides:
        t = s.notes_slide.notes_text_frame
        if t is not None:
            src_sp = t._txBody.getparent()
            break
    assert src_sp is not None, "no working notes placeholder found"

    # ---- p2 这几天工作总览
    wipe(s_over)
    header(s_over, "进展 · 01", "这几天工作总览", dk.TEAL)
    section_label(s_over, 0.9, 1.95, "完成的工作", dk.TEAL)
    for i, t in enumerate(OVERVIEW_WORK):
        numbered_row(s_over, 0.9, 2.4 + i * 0.78, 6.2, 0.66, i + 1, t, dk.TEAL)
    section_label(s_over, 7.4, 1.95, "形成的产出", dk.CORAL)
    for i, t in enumerate(OVERVIEW_OUT):
        output_row(s_over, 7.4, 2.4 + i * 1.0, 5.03, 0.84, t, dk.CORAL)
    band(s_over, "5 项工作已完成，4 项阶段 0 产出已就位", dk.TEAL)
    set_notes(s_over, NOTES[1], src_sp)

    # ---- p3 需求与信号链梳理
    wipe(s_arch)
    header(s_arch, "进展 · 02", "需求与信号链梳理", dk.CORAL)
    white_card(s_arch, 0.9, 1.9, 6.0, 2.85)
    section_label(s_arch, 1.15, 2.08, "需求理解", dk.CORAL)
    bullet_block(s_arch, 1.15, 2.5, 5.55, 2.1, NEEDS, size=15.5, gap=6)
    white_card(s_arch, 7.2, 1.9, 5.23, 2.85)
    section_label(s_arch, 7.45, 2.08, "关键说明", dk.AMBER)
    bullet_block(s_arch, 7.45, 2.5, 4.75, 2.1, NOTES_TXT, size=15.5, gap=6)
    section_label(s_arch, 0.9, 4.85, "信号链", dk.TEAL)
    x = 0.9
    for i, (label, w, fill) in enumerate(CHAIN):
        rect(s_arch, x, 5.18, w, 0.86, fill=fill, rounded=True, radius=0.14,
             text=[{"text": label, "size": 13, "bold": True, "color": dk.WHITE,
                    "align": "c", "line_spacing": 1.15}],
             anchor="MIDDLE", pad=(0.06, 0.03, 0.06, 0.03))
        x += w
        if i < len(CHAIN) - 1:
            txt(s_arch, x, 5.18, 0.16, 0.86,
                [{"text": "›", "size": 15, "bold": True, "color": dk.MUTED,
                  "align": "c"}], anchor="MIDDLE")
            x += 0.16
    band(s_arch, "需求已明确：脱机可完成本地诊断，服务器只做外部存储", dk.CORAL)
    set_notes(s_arch, NOTES[2], src_sp)

    # ---- p4 数据接口与协议初稿
    wipe(s_duty)
    header(s_duty, "进展 · 03", "数据接口与协议初稿", dk.AMBER)
    section_label(s_duty, 0.9, 1.9, "帧结构（每帧 352 B）", dk.AMBER)
    bar_y, bar_h = 2.3, 0.72
    segs = [(0.9, 8.6, dk.TEAL, "payload · 320 B", 14),
            (9.5, 2.0, dk.INK, "header · 30 B", 13),
            (11.5, 0.93, dk.CORAL, "CRC 2 B", 11)]
    for sx, sw, sc, st, ss in segs:
        rect(s_duty, sx, bar_y, sw, bar_h, fill=sc,
             text=[{"text": st, "size": ss, "bold": True, "color": dk.WHITE,
                    "align": "c"}],
             anchor="MIDDLE", pad=(0.05, 0.02, 0.05, 0.02))
    txt(s_duty, 0.9, 3.12, 11.53, 0.3,
        [{"text": "20 ms 一帧 · 160 samples @ 8 kHz · 16-bit 单声道小端",
          "size": 13.5, "color": dk.MUTED}])
    rect(s_duty, 6.72, 3.62, 0.015, 2.6, fill=dk.CARD_LINE)
    section_label(s_duty, 0.9, 3.58, "采集与帧格式", dk.AMBER)
    bullet_block(s_duty, 0.9, 4.0, 5.6, 2.3, IF_LEFT, size=16, gap=8)
    section_label(s_duty, 6.95, 3.58, "时间与责任", dk.AMBER)
    bullet_block(s_duty, 6.95, 4.0, 5.48, 2.3, IF_RIGHT, size=15.5, gap=8)
    band(s_duty, "接口初稿成型：8 kHz 采集 · 20 ms 帧 · 352 B/帧，序号时间戳由探头生成",
         dk.AMBER)
    set_notes(s_duty, NOTES[3], src_sp)

    # ---- p7 需要老师确认：rebuild from outline v3 (the user's own 11:07
    # in-deck edits dropped rows the outline still requires; outline-3 is newer)
    wipe(s_conf)
    header(s_conf, "待确认 · 01", "需要老师确认：结构大方向与上传范围", dk.CORAL)
    cw7 = (dk.CW - 0.4) / 2
    for i, (head, accent, rows) in enumerate(CONFIRM):
        confirm_card(s_conf, 0.9 + i * (cw7 + 0.4), cw7, accent, head, rows)
    band(s_conf, "请老师重点确认：总体结构 · 探头与主机连接方式 · 上传数据范围",
         dk.CORAL)
    set_notes(s_conf, NOTES[6], src_sp)

    # ---- p6 团队分工 kicker
    replace_text(s_team, "进展 · 04", "进展 · 05")
    set_notes(s_team, NOTES[5], src_sp)

    # ---- p5 系统架构与产品方案 (create once; reuse on re-runs)
    blank = prs.slide_layouts[6]
    try:
        s_new = find_slide(prs, "系统架构与产品方案")
        wipe(s_new)
    except SystemExit:
        s_new = prs.slides.add_slide(blank)
    header(s_new, "进展 · 04", "系统架构与产品方案", dk.TEAL)
    flow_strip(s_new, top=1.85, height=0.9)
    section_label(s_new, 0.9, 2.88, "系统架构", dk.TEAL)
    bullet_block(s_new, 0.9, 3.22, 11.53, 1.2, ARCH_BULLETS, size=16, gap=5)
    section_label(s_new, 0.9, 4.5, "产品方案", dk.CORAL)
    bullet_block(s_new, 0.9, 4.84, 5.7, 1.45, PRODUCT_BULLETS, size=15, gap=5)
    white_card(s_new, 6.85, 4.42, 5.58, 1.9)
    img(s_new, ARCH, 7.28, 4.56, 2.62, 1.6)
    img(s_new, PRODUCT, 10.15, 4.56, 1.84, 1.6)
    band(s_new, "二体式已定：探头采集、主机本地推理与显示、服务器只存储", dk.TEAL)
    set_notes(s_new, NOTES[4], src_sp)

    # ---- p8 下一步 (create once; reuse on re-runs)
    try:
        s_next = find_slide(prs, "计划 · 01")
        wipe(s_next)
    except SystemExit:
        s_next = prs.slides.add_slide(blank)
    header(s_next, "计划 · 01", "下一步", dk.TEAL)
    cw = (dk.CW - 2 * 0.3) / 3
    for i, (num, head, body, accent) in enumerate(STEPS):
        x = 0.9 + i * (cw + 0.3)
        y, h = 2.3, 3.7
        white_card(s_next, x, y, cw, h)
        rect(s_next, x, y, cw, 0.1, fill=accent, rounded=True, radius=0.4)
        rect(s_next, x + 0.3, y + 0.5, 0.62, 0.62, fill=accent,
             rounded=True, radius=0.5,
             text=[{"text": num, "size": 22, "bold": True, "color": dk.WHITE,
                    "align": "c"}],
             anchor="MIDDLE", pad=(0, 0, 0, 0))
        txt(s_next, x + 0.3, y + 1.48, cw - 0.6, 0.85,
            [{"text": head, "size": 20, "bold": True, "color": dk.INK,
              "line_spacing": 1.2}])
        txt(s_next, x + 0.3, y + 2.45, cw - 0.6, 1.15,
            [{"text": body, "size": 15.5, "color": dk.BODY,
              "line_spacing": 1.3}])
    band(s_next, "顺序：先冻结结构与连接 → 再定上传范围 → 最后细化器件、接口与上传字段",
         dk.TEAL)
    set_notes(s_next, NOTES[7], src_sp)

    # ---- order: move 系统架构与产品方案 before 团队分工
    lst = prs.slides._sldIdLst
    moving = None
    for el in list(lst):
        sld = prs.part.related_part(el.get(qn("r:id"))).slide
        if any(sh.has_text_frame and "系统架构与产品方案" in sh.text_frame.text
               for sh in sld.shapes):
            moving = el
    assert moving is not None
    lst.remove(moving)
    for i, el in enumerate(list(lst)):
        sld = prs.part.related_part(el.get(qn("r:id"))).slide
        if any(sh.has_text_frame and "我负责" in sh.text_frame.text
               for sh in sld.shapes):
            lst.insert(i, moving)
            break
    else:
        raise SystemExit("团队分工 slide not found for reorder")

    # ---- explicit backgrounds (cover dark, pages cream)
    for i, slide in enumerate(prs.slides):
        marker = "FINAL YEAR PROJECT"
        is_cover = any(sh.has_text_frame and marker in sh.text_frame.text
                       for sh in slide.shapes)
        set_bg(slide, dk.INK if is_cover else dk.BASE)

    # ---- animation groups (regions over final geometry)
    regions = {
        1: [(0.9, 1.85, 7.1, 6.4), (7.1, 1.85, 13.4, 6.4)],       # 工作 / 产出
        2: [(0.0, 1.75, 7.1, 4.7), (7.1, 1.75, 13.4, 4.7),
            (0.0, 4.7, 13.4, 6.4)],                                 # 需求 / 说明 / 链
        3: [(0.0, 1.75, 13.4, 3.45), (0.0, 3.45, 6.7, 6.4),
            (6.7, 3.45, 13.4, 6.4)],                                # 帧结构 / 左 / 右
        4: [(0.0, 1.75, 13.4, 2.82), (0.0, 2.82, 13.4, 4.35),
            (0.0, 4.35, 13.4, 6.4)],                                # 流程 / 架构 / 产品+图
        5: [(0.0, 1.75, 4.7, 6.4), (4.7, 1.75, 8.65, 6.4),
            (8.65, 1.75, 13.4, 6.4)],                               # 团队三列
        6: [(0.0, 1.75, 6.7, 6.4), (6.7, 1.75, 13.4, 6.4)],         # 确认两卡
        7: [(0.0, 1.75, 4.7, 6.4), (4.7, 1.75, 8.65, 6.4),
            (8.65, 1.75, 13.4, 6.4)],                               # 下一步三步
    }
    groups = {}
    for idx, slide in enumerate(prs.slides):
        if idx not in regions:
            continue
        rects = []
        for sh in slide.shapes:
            cx = (sh.left + sh.width / 2) / 914400.0
            cy = (sh.top + sh.height / 2) / 914400.0
            if cy < 1.75 or cy >= 6.4:
                continue
            rects.append((sh.shape_id, cx, cy))
        gg = [[i for i, cx, cy in rects if x0 <= cx < x1 and y0 <= cy < y1]
              for x0, y0, x1, y1 in regions[idx]]
        gg = [g for g in gg if g]
        if gg:
            groups[str(idx)] = gg
    json.dump(groups, open(GROUPS_OUT, "w"))
    print("animation groups ->", GROUPS_OUT)
    for k in sorted(groups, key=int):
        print("  p%d" % (int(k) + 1), [len(g) for g in groups[k]])

    try:
        prs.save(DECK)
        print("saved in place:", DECK)
    except PermissionError:
        alt = Path(str(DECK).replace(".pptx", "-v3.pptx"))
        prs.save(alt)
        print("PERMISSION ERROR -> saved to", alt)


if __name__ == "__main__":
    build()
