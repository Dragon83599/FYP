# -*- coding: utf-8 -*-
"""Post-revise visual fix pass (in place, content-addressed by geometry):

1. Slide 3 statement rows used TINT_ROW (#F7F4EF) directly on the cream page
   background (#F5F2EC) -> they were nearly invisible and the left column
   read as empty. Give them a white card fill + hairline border.
2. Slide 6 numbered rows sit on white cards; strengthen the tint so the rows
   separate from the card instead of looking blank.
"""
import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[3]
DECK = ROOT / "meetings" / "2026-10-06" / "group-meeting-2026-10-06.pptx"
ROW_TINT = "F1EBE1"


def inches(v):
    return v / 914400.0


def main():
    prs = Presentation(DECK)
    slides = list(prs.slides)

    # --- slide 3 (index 2): statement rows -> white cards w/ border
    n3 = 0
    for sh in slides[2].shapes:
        l, t = inches(sh.left), inches(sh.top)
        w, h = inches(sh.width), inches(sh.height)
        if abs(l - 0.9) < 0.02 and abs(w - 5.5) < 0.02 and abs(h - 0.96) < 0.02:
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
            sh.line.color.rgb = RGBColor.from_string("E4DED4")
            sh.line.width = Pt(1.0)
            n3 += 1
    print("slide 3 rows restyled:", n3)

    # --- slide 6 (index 5): numbered rows -> stronger tint
    n6 = 0
    for sh in slides[5].shapes:
        l, t = inches(sh.left), inches(sh.top)
        w, h = inches(sh.width), inches(sh.height)
        if abs(h - 0.74) < 0.02 and abs(w - 4.97) < 0.03 and t > 2.6:
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string(ROW_TINT)
            n6 += 1
    print("slide 6 rows retinted:", n6)

    assert n3 == 3 and n6 == 8, (n3, n6)
    prs.save(DECK)
    print("saved:", DECK)


if __name__ == "__main__":
    main()
