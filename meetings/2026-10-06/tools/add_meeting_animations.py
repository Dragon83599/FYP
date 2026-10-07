# -*- coding: utf-8 -*-
"""Add per-click fade entrance animations + 0.7s fade transitions to
group-meeting-2026-10-06.pptx.

Shape ids are read back from the FINAL file (never assumed), and the script
is idempotent: if a slide already carries <p:timing>, it is rebuilt from
scratch rather than appended to.
"""
import json
import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml import parse_xml

ROOT = Path(__file__).resolve().parents[3]
DECK = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "meetings" / "2026-10-06" / "group-meeting-2026-10-06.pptx"
TMP = Path(os.environ.get("LOCALAPPDATA", r"C:\Temp")) / "Temp"
GROUPS = os.path.join(TMP, "deck_anim_groups.json")

P = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
P14 = 'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"'

TRANSITION = (
    f'<mc:AlternateContent {MC} {P} {P14}>'
    '<mc:Choice Requires="p14">'
    '<p:transition spd="med" p14:dur="700"><p:fade/></p:transition>'
    '</mc:Choice>'
    '<mc:Fallback>'
    f'<p:transition {P} spd="med"><p:fade/></p:transition>'
    '</mc:Fallback>'
    '</mc:AlternateContent>'
)


def click_block(state, spids):
    """One click = one level-1 par; every shape in it fades in together."""
    cid = state.next()
    lvl2 = state.next()
    effects = []
    for k, spid in enumerate(spids):
        eid, sid, aid = state.next(), state.next(), state.next()
        node = "clickEffect" if k == 0 else "withEffect"
        effects.append(
            f'<p:par><p:cTn id="{eid}" presetID="10" presetClass="entr" '
            f'presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            '<p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{sid}" dur="1" fill="hold">'
            '<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
            '<p:attrNameLst><p:attrName>style.visibility</p:attrName>'
            '</p:attrNameLst></p:cBhvr>'
            '<p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade">'
            f'<p:cBhvr><p:cTn id="{aid}" dur="500">'
            '<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
            '</p:cBhvr></p:animEffect>'
            '</p:childTnLst></p:cTn></p:par>')
    return (
        f'<p:par><p:cTn id="{cid}" fill="hold">'
        '<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
        f'<p:childTnLst><p:par><p:cTn id="{lvl2}" fill="hold">'
        '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f'<p:childTnLst>{"".join(effects)}</p:childTnLst>'
        '</p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


class State:
    def __init__(self):
        self.n = 2  # 1 = tmRoot, 2 = mainSeq

    def next(self):
        self.n += 1
        return self.n


def timing_xml(groups):
    st = State()
    blocks = "".join(click_block(st, g) for g in groups)
    return (
        f'<p:timing {P}><p:tnLst>'
        '<p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        '<p:childTnLst><p:seq concurrent="1" nextAc="seek">'
        '<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'{blocks}'
        '</p:childTnLst></p:cTn>'
        '<p:prevCondLst><p:cond evt="onPrev" delay="0">'
        '<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0">'
        '<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        '</p:seq></p:childTnLst></p:cTn></p:par>'
        '</p:tnLst></p:timing>')


def main():
    groups_by_slide = {int(k): v for k, v in json.load(open(GROUPS)).items()}
    prs = Presentation(DECK)
    slides = list(prs.slides)
    known = set()
    for si, slide in enumerate(slides):
        root = slide._element
        # idempotency: drop any previous transition/timing
        for tag in ("{http://schemas.openxmlformats.org/presentationml/2006/main}transition",
                    "{http://schemas.openxmlformats.org/presentationml/2006/main}timing"):
            for el in root.findall(tag):
                root.remove(el)
        for el in root.findall(
                "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"):
            root.remove(el)

        shape_ids = {sh.shape_id for sh in slide.shapes}
        groups = groups_by_slide.get(si, [])
        # validate every spid exists on this slide
        checked = []
        for g in groups:
            ids = [i for i in g if i in shape_ids]
            missing = set(g) - set(ids)
            if missing:
                print(f"  slide {si + 1}: MISSING spids {missing} -> skipped")
            if ids:
                checked.append(ids)
        if checked:
            known |= {i for g in checked for i in g}
            n_clicks = len(checked)
        else:
            n_clicks = 0
        # p:sld child order must be cSld, clrMapOvr, transition, timing
        root.append(parse_xml(TRANSITION))
        if checked:
            root.append(parse_xml(timing_xml(checked)))
        print(f"slide {si + 1}: {n_clicks} click group(s), transition fade 700ms")

    prs.save(DECK)

    # verify: reload and assert spids referenced exist, and timing present
    prs2 = Presentation(DECK)
    ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    total = 0
    for si, slide in enumerate(prs2.slides, 1):
        ids = {sh.shape_id for sh in slide.shapes}
        timing = slide._element.find(ns + "timing")
        if timing is None:
            continue
        tgt = [int(e.get("spid")) for e in timing.iter(ns + "spTgt")]
        bad = [t for t in tgt if t not in ids]
        assert not bad, f"slide {si}: spid not on slide: {bad}"
        total += len(set(tgt))
        trans = (slide._element.find(ns + "transition") is not None
                 or slide._element.find(
                     "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")
                 is not None)
        assert trans, f"slide {si}: no transition"
    print(f"OK: {total} animated shapes across all slides; file re-opened cleanly")


if __name__ == "__main__":
    sys.exit(main())
